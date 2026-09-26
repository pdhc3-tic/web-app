"""Exportações do SGP: fluxo assíncrono (`/sgp/exportacoes/`), exportação de
Atividades e exportação de UPFs com os filtros da listagem."""
import csv
import io
from datetime import datetime, timedelta
from unittest.mock import patch

import pytest
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APIClient

from apps.core.tests.factories import RoleFactory, UserFactory
from apps.sgp.models import UPF, ExportJob
from apps.sgp.services import exportacao as exportacao_service
from apps.sgp.services import upf_export
from apps.sgp.tests.factories import (
    ActivityFactory,
    ComunidadeFactory,
    UPFFactory,
    WorkPlanAcaoFactory,
)

pytestmark = pytest.mark.django_db

EXPORTACOES_URL = "/api/v1/sgp/exportacoes/"
ATIVIDADES_EXPORT_URL = "/api/v1/sgp/atividades/exportar/"
UPFS_URL = "/api/v1/upfs/"
UPFS_EXPORT_URL = "/api/v1/upfs/exportar/"
DELAY = "apps.sgp.tasks.processar_exportacao.delay"


def _linhas_csv(conteudo: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(conteudo.decode("utf-8-sig"))))


def _detalhe(job_id):
    return f"{EXPORTACOES_URL}{job_id}/"


def _cliente(user):
    # As fixtures de cliente autenticado compartilham uma única instância de
    # APIClient; um segundo perfil no mesmo teste precisa da sua.
    cliente = APIClient()
    cliente.force_authenticate(user=user)
    return cliente


@pytest.fixture
def usuario_ugp(db):
    return UserFactory(profiles=[(RoleFactory(slug="ugp", nome="UGP"), None)])


@pytest.fixture
def cliente_ugp(api_client, usuario_ugp):
    api_client.force_authenticate(user=usuario_ugp)
    return api_client


def _criar_exportacao(cliente, django_capture_on_commit_callbacks, **payload):
    with patch(DELAY) as delay:
        with django_capture_on_commit_callbacks(execute=True):
            response = cliente.post(EXPORTACOES_URL, data=payload, format="json")
    return response, delay


# ===========================================================================
# Fluxo assíncrono
# ===========================================================================

class TestFluxoAssincrono:
    def test_criar_retorna_202_e_enfileira(self, cliente_ugp, django_capture_on_commit_callbacks):
        response, delay = _criar_exportacao(
            cliente_ugp, django_capture_on_commit_callbacks,
            tipo="atividades", formato="csv",
            filtros={"periodo_inicio": "2026-01-01", "periodo_fim": "2026-12-31"},
        )

        assert response.status_code == status.HTTP_202_ACCEPTED, response.data
        assert response.data["status"] == "pendente"
        delay.assert_called_once_with(response.data["id"])
        job = ExportJob.objects.get(pk=response.data["id"])
        assert job.filtros == {"periodo_inicio": "2026-01-01", "periodo_fim": "2026-12-31"}

    def test_status_download_apos_processar(self, cliente_ugp, municipio_rn, django_capture_on_commit_callbacks):
        ActivityFactory(titulo="Oficina de caprinos", municipio=municipio_rn)
        response, _ = _criar_exportacao(
            cliente_ugp, django_capture_on_commit_callbacks, tipo="atividades", formato="csv",
        )
        job_id = response.data["id"]

        exportacao_service.executar_exportacao(job_id)

        detalhe = cliente_ugp.get(_detalhe(job_id))
        assert detalhe.data["status"] == "concluida"
        assert detalhe.data["progresso"] == 100
        assert detalhe.data["total_registros"] == 1
        assert detalhe.data["expira_em"] is not None

        download = cliente_ugp.get(f"{_detalhe(job_id)}download/")
        assert download.status_code == status.HTTP_200_OK
        assert download["Content-Type"] == "text/csv; charset=utf-8"
        assert download["Content-Disposition"].endswith('.csv"')
        linhas = _linhas_csv(download.content)
        assert linhas[0][:2] == ["ID", "Título"]
        assert linhas[1][1] == "Oficina de caprinos"

    def test_xlsx(self, cliente_ugp, django_capture_on_commit_callbacks):
        from openpyxl import load_workbook

        ActivityFactory()
        response, _ = _criar_exportacao(
            cliente_ugp, django_capture_on_commit_callbacks, tipo="atividades", formato="xlsx",
        )
        exportacao_service.executar_exportacao(response.data["id"])

        download = cliente_ugp.get(f"{_detalhe(response.data['id'])}download/")

        workbook = load_workbook(io.BytesIO(download.content))
        assert workbook.sheetnames == ["Atividades"]

    def test_download_antes_de_concluir_retorna_409(self, cliente_ugp, django_capture_on_commit_callbacks):
        response, _ = _criar_exportacao(
            cliente_ugp, django_capture_on_commit_callbacks, tipo="plano_trabalho", formato="csv",
        )

        download = cliente_ugp.get(f"{_detalhe(response.data['id'])}download/")

        assert download.status_code == status.HTTP_409_CONFLICT
        assert download.data["code"] == "exportacao_nao_concluida"
        assert download.data["status"] == "pendente"

    def test_download_expirado_retorna_410(self, cliente_ugp, usuario_ugp):
        job = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp,
            status=ExportJob.Status.CONCLUIDA, conteudo=b"x", nome_arquivo="a.csv",
            content_type="text/csv", expira_em=timezone.now() - timedelta(minutes=1),
        )

        download = cliente_ugp.get(f"{_detalhe(job.pk)}download/")

        assert download.status_code == status.HTTP_410_GONE
        assert download.data["code"] == "exportacao_expirada"

    def test_so_o_dono_acessa(self, usuario_ugp):
        job = ExportJob.objects.create(tipo="atividades", formato="csv", solicitante=usuario_ugp)
        outro = _cliente(UserFactory(profiles=[(RoleFactory(slug="ugp", nome="UGP"), None)]))

        assert outro.get(_detalhe(job.pk)).status_code == status.HTTP_404_NOT_FOUND
        assert outro.get(f"{_detalhe(job.pk)}download/").status_code == status.HTTP_404_NOT_FOUND
        assert outro.post(f"{_detalhe(job.pk)}repetir/").status_code == status.HTTP_404_NOT_FOUND

    def test_filtro_invalido_termina_em_erro_e_pode_repetir(
        self, cliente_ugp, usuario_ugp, django_capture_on_commit_callbacks
    ):
        job = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp,
            filtros={"territorio_id": "abc"},
        )

        exportacao_service.executar_exportacao(job.pk)

        job.refresh_from_db()
        assert job.status == ExportJob.Status.ERRO
        assert "territorio_id" in job.erro

        with patch(DELAY) as delay:
            with django_capture_on_commit_callbacks(execute=True):
                response = cliente_ugp.post(f"{_detalhe(job.pk)}repetir/")
        assert response.status_code == status.HTTP_202_ACCEPTED
        assert response.data["status"] == "pendente"
        assert response.data["erro"] == ""
        delay.assert_called_once_with(job.pk)

    def test_falha_inesperada_vira_erro_generico(self, usuario_ugp):
        job = ExportJob.objects.create(tipo="atividades", formato="csv", solicitante=usuario_ugp)

        with patch.object(exportacao_service, "gerar_arquivo", side_effect=RuntimeError("disco")):
            exportacao_service.executar_exportacao(job.pk)

        job.refresh_from_db()
        assert job.status == ExportJob.Status.ERRO
        assert "disco" not in job.erro
        assert job.erro.startswith("Falha inesperada")

    def test_solicitante_que_perdeu_acesso_recebe_erro(self, db):
        sem_perfil = UserFactory()
        job = ExportJob.objects.create(tipo="atividades", formato="csv", solicitante=sem_perfil)

        exportacao_service.executar_exportacao(job.pk)

        job.refresh_from_db()
        assert job.status == ExportJob.Status.ERRO
        assert "acesso" in job.erro

    def test_criacao_registra_enfileirado_em(self, cliente_ugp, django_capture_on_commit_callbacks):
        response, _ = _criar_exportacao(
            cliente_ugp, django_capture_on_commit_callbacks, tipo="atividades", formato="csv",
        )

        assert ExportJob.objects.get(pk=response.data["id"]).enfileirado_em is not None

    @pytest.mark.parametrize("status_job,campo", [
        (ExportJob.Status.PENDENTE, "enfileirado_em"),
        (ExportJob.Status.PROCESSANDO, "iniciado_em"),
    ])
    def test_job_travado_pode_ser_repetido(
        self, cliente_ugp, usuario_ugp, django_capture_on_commit_callbacks, status_job, campo
    ):
        antigo = timezone.now() - exportacao_service.PRAZO_JOB_TRAVADO - timedelta(minutes=1)
        job = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp,
            status=status_job, **{campo: antigo},
        )

        with patch(DELAY) as delay:
            with django_capture_on_commit_callbacks(execute=True):
                response = cliente_ugp.post(f"{_detalhe(job.pk)}repetir/")

        assert response.status_code == status.HTTP_202_ACCEPTED
        job.refresh_from_db()
        assert job.status == ExportJob.Status.PENDENTE
        assert job.enfileirado_em > antigo
        delay.assert_called_once_with(job.pk)

    @pytest.mark.parametrize("status_job,campo", [
        (ExportJob.Status.PENDENTE, "enfileirado_em"),
        (ExportJob.Status.PROCESSANDO, "iniciado_em"),
    ])
    def test_job_em_andamento_dentro_do_prazo_nao_pode_ser_repetido(
        self, cliente_ugp, usuario_ugp, status_job, campo
    ):
        job = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp,
            status=status_job, **{campo: timezone.now()},
        )

        response = cliente_ugp.post(f"{_detalhe(job.pk)}repetir/")

        assert response.status_code == status.HTTP_409_CONFLICT
        assert response.data["code"] == "exportacao_nao_repetivel"

    def test_repetir_so_com_erro(self, cliente_ugp, usuario_ugp):
        job = ExportJob.objects.create(tipo="atividades", formato="csv", solicitante=usuario_ugp)

        response = cliente_ugp.post(f"{_detalhe(job.pk)}repetir/")

        assert response.status_code == status.HTTP_409_CONFLICT
        assert response.data["code"] == "exportacao_nao_repetivel"

    def test_reentrega_do_broker_nao_reprocessa(self, usuario_ugp):
        job = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp,
            status=ExportJob.Status.CONCLUIDA, conteudo=b"original",
        )

        exportacao_service.executar_exportacao(job.pk)

        job.refresh_from_db()
        assert bytes(job.conteudo) == b"original"

    def test_validacao_na_criacao(self, cliente_ugp):
        tipo_invalido = cliente_ugp.post(
            EXPORTACOES_URL, data={"tipo": "membros", "formato": "csv"}, format="json"
        )
        periodo_invertido = cliente_ugp.post(
            EXPORTACOES_URL,
            data={
                "tipo": "atividades", "formato": "csv",
                "filtros": {"periodo_inicio": "2026-12-01", "periodo_fim": "2026-01-01"},
            },
            format="json",
        )
        parametro_upf = cliente_ugp.post(
            EXPORTACOES_URL,
            data={"tipo": "upfs", "formato": "csv", "filtros": {"situacao": "ativas"}},
            format="json",
        )

        assert tipo_invalido.status_code == status.HTTP_400_BAD_REQUEST
        assert periodo_invertido.status_code == status.HTTP_400_BAD_REQUEST
        assert "filtros" in periodo_invertido.data
        assert parametro_upf.status_code == status.HTTP_400_BAD_REQUEST
        assert parametro_upf.data["code"] == "parametro_desconhecido"
        assert ExportJob.objects.count() == 0

    def test_valor_invalido_de_filtro_de_upf_e_recusado_na_criacao(self, cliente_ugp):
        response = cliente_ugp.post(
            EXPORTACOES_URL,
            data={"tipo": "upfs", "formato": "csv", "filtros": {"municipio": "abc"}},
            format="json",
        )

        assert response.status_code == status.HTTP_400_BAD_REQUEST
        assert "municipio" in response.data["filtros"]
        assert ExportJob.objects.count() == 0

    def test_usuario_sem_perfil_sgp_recebe_403(self, auth_client_sem_acesso):
        response = auth_client_sem_acesso.post(
            EXPORTACOES_URL, data={"tipo": "atividades", "formato": "csv"}, format="json"
        )

        assert response.status_code == status.HTTP_403_FORBIDDEN

    @pytest.mark.parametrize("tipo,filtros", [
        ("atividades", {"territorio": "1"}),
        ("atividades", {"meta_id": "1"}),
        ("plano_trabalho", {"acao_id": "1"}),
        ("plano_trabalho", {"formato": "xlsx"}),
    ])
    def test_filtro_desconhecido_e_recusado_em_qualquer_tipo(self, cliente_ugp, tipo, filtros):
        response = cliente_ugp.post(
            EXPORTACOES_URL, data={"tipo": tipo, "formato": "csv", "filtros": filtros}, format="json"
        )

        assert response.status_code == status.HTTP_400_BAD_REQUEST
        assert response.data["code"] == "parametro_desconhecido"
        assert response.data["parametros"] == sorted(filtros)
        assert ExportJob.objects.count() == 0

    def test_job_respeita_o_escopo_do_solicitante(self, usuario_adt_rn, municipio_rn, municipio_ce):
        do_rn = ActivityFactory(municipio=municipio_rn)
        ActivityFactory(municipio=municipio_ce)
        UPFFactory(municipio=municipio_rn, _titular_nome="Do RN")
        UPFFactory(municipio=municipio_ce, _titular_nome="Do CE")
        atividades = ExportJob.objects.create(tipo="atividades", formato="csv", solicitante=usuario_adt_rn)
        upfs = ExportJob.objects.create(tipo="upfs", formato="csv", solicitante=usuario_adt_rn)

        exportacao_service.executar_exportacao(atividades.pk)
        exportacao_service.executar_exportacao(upfs.pk)

        atividades.refresh_from_db()
        upfs.refresh_from_db()
        assert [linha[0] for linha in _linhas_csv(bytes(atividades.conteudo))[1:]] == [str(do_rn.pk)]
        linhas_upfs = _linhas_csv(bytes(upfs.conteudo))
        coluna_titular = linhas_upfs[0].index("Titular")
        assert [linha[coluna_titular] for linha in linhas_upfs[1:]] == ["Do RN"]


class TestJobsTravados:
    def test_marca_como_erro_so_os_parados_alem_do_prazo(self, usuario_ugp):
        agora = timezone.now()
        antigo = agora - exportacao_service.PRAZO_JOB_TRAVADO - timedelta(minutes=1)

        def job(**campos):
            return ExportJob.objects.create(
                tipo="atividades", formato="csv", solicitante=usuario_ugp, **campos
            )

        pendente_perdido = job(status=ExportJob.Status.PENDENTE, enfileirado_em=antigo)
        worker_morto = job(status=ExportJob.Status.PROCESSANDO, iniciado_em=antigo)
        pendente_recente = job(status=ExportJob.Status.PENDENTE, enfileirado_em=agora)
        processando_recente = job(status=ExportJob.Status.PROCESSANDO, iniciado_em=agora)
        concluida_antiga = job(status=ExportJob.Status.CONCLUIDA, iniciado_em=antigo)

        assert exportacao_service.marcar_exportacoes_travadas() == 2

        status_por_job = dict(ExportJob.objects.values_list("pk", "status"))
        assert status_por_job[pendente_perdido.pk] == ExportJob.Status.ERRO
        assert status_por_job[worker_morto.pk] == ExportJob.Status.ERRO
        assert status_por_job[pendente_recente.pk] == ExportJob.Status.PENDENTE
        assert status_por_job[processando_recente.pk] == ExportJob.Status.PROCESSANDO
        assert status_por_job[concluida_antiga.pk] == ExportJob.Status.CONCLUIDA
        worker_morto.refresh_from_db()
        assert worker_morto.erro == exportacao_service.MENSAGEM_JOB_TRAVADO

    def test_task_e_interrompida_antes_de_o_job_contar_como_travado(self):
        from apps.sgp.tasks import processar_exportacao

        prazo = exportacao_service.PRAZO_JOB_TRAVADO.total_seconds()
        assert processar_exportacao.soft_time_limit < processar_exportacao.time_limit < prazo


class TestLimpeza:
    def test_remove_expiradas_e_abandonadas(self, usuario_ugp):
        agora = timezone.now()
        expirada = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp,
            status=ExportJob.Status.CONCLUIDA, expira_em=agora - timedelta(hours=1),
        )
        valida = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp,
            status=ExportJob.Status.CONCLUIDA, expira_em=agora + timedelta(hours=1),
        )
        abandonada = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp, status=ExportJob.Status.ERRO,
        )
        ExportJob.objects.filter(pk=abandonada.pk).update(criado_em=agora - timedelta(days=8))
        recente_com_erro = ExportJob.objects.create(
            tipo="atividades", formato="csv", solicitante=usuario_ugp, status=ExportJob.Status.ERRO,
        )

        removidas = exportacao_service.limpar_exportacoes_expiradas()

        assert removidas == 2
        restantes = set(ExportJob.objects.values_list("pk", flat=True))
        assert restantes == {valida.pk, recente_com_erro.pk}
        assert expirada.pk not in restantes


# ===========================================================================
# Exportação de Atividades (download direto)
# ===========================================================================

class TestExportacaoAtividades:
    def test_filtros_e_colunas(self, cliente_ugp, municipio_rn, municipio_ce):
        acao = WorkPlanAcaoFactory()
        dentro = ActivityFactory(
            titulo="Dentro", municipio=municipio_rn, acao=acao,
            data_inicio=timezone.make_aware(datetime(2026, 3, 10, 8)),
            data_fim=timezone.make_aware(datetime(2026, 3, 10, 12)),
        )
        ActivityFactory(  # fora do período
            municipio=municipio_rn, acao=acao,
            data_inicio=timezone.make_aware(datetime(2025, 3, 10, 8)),
            data_fim=timezone.make_aware(datetime(2025, 3, 10, 12)),
        )
        ActivityFactory(municipio=municipio_ce, acao=acao)  # outro território
        ActivityFactory(municipio=municipio_rn)  # outra ação

        response = cliente_ugp.get(ATIVIDADES_EXPORT_URL, {
            "formato": "csv",
            "periodo_inicio": "2026-01-01",
            "periodo_fim": "2026-12-31",
            "territorio_id": municipio_rn.territory_id,
            "acao_id": acao.pk,
        })

        assert response.status_code == status.HTTP_200_OK
        linhas = _linhas_csv(response.content)
        assert len(linhas) == 2
        cabecalho = linhas[0]
        registro = dict(zip(cabecalho, linhas[1]))
        assert registro["ID"] == str(dentro.pk)
        assert registro["Estado"] == "RN"
        assert registro["Município"] == municipio_rn.nome
        assert registro["Território"] == municipio_rn.territory.nome
        assert registro["Ação"].startswith(acao.numero)

    def test_contagem_de_participantes(self, cliente_ugp, municipio_rn):
        atividade = ActivityFactory(municipio=municipio_rn)
        upfs = UPFFactory.create_batch(2, municipio=municipio_rn)
        atividade.upfs_participantes.set(upfs)
        atividade.membros_participantes.set([upf.titular for upf in upfs])

        response = cliente_ugp.get(ATIVIDADES_EXPORT_URL, {"formato": "csv"})

        registro = dict(zip(*_linhas_csv(response.content)[:2]))
        assert registro["UPFs participantes"] == "2"
        assert registro["Participantes"] == "2"

    def test_escopo_territorial(self, auth_client_adt_rn, municipio_rn, municipio_ce):
        do_rn = ActivityFactory(municipio=municipio_rn)
        ActivityFactory(municipio=municipio_ce)

        response = auth_client_adt_rn.get(ATIVIDADES_EXPORT_URL, {"formato": "csv"})

        ids = [linha[0] for linha in _linhas_csv(response.content)[1:]]
        assert ids == [str(do_rn.pk)]

    def test_periodo_invertido_retorna_400(self, cliente_ugp):
        response = cliente_ugp.get(ATIVIDADES_EXPORT_URL, {
            "formato": "csv", "periodo_inicio": "2026-12-01", "periodo_fim": "2026-01-01",
        })

        assert response.status_code == status.HTTP_400_BAD_REQUEST

    def test_formato_obrigatorio(self, cliente_ugp):
        assert cliente_ugp.get(ATIVIDADES_EXPORT_URL).status_code == status.HTTP_400_BAD_REQUEST

    def test_xlsx(self, cliente_ugp):
        response = cliente_ugp.get(ATIVIDADES_EXPORT_URL, {"formato": "xlsx"})

        assert response.status_code == status.HTTP_200_OK
        assert response["Content-Type"].startswith("application/vnd.openxmlformats")

    def test_usuario_sem_perfil_sgp_recebe_403(self, auth_client_sem_acesso):
        response = auth_client_sem_acesso.get(ATIVIDADES_EXPORT_URL, {"formato": "csv"})

        assert response.status_code == status.HTTP_403_FORBIDDEN


# ===========================================================================
# Exportação de UPFs
# ===========================================================================

class TestExportacaoUPFs:
    def _ids_listagem(self, cliente, params):
        response = cliente.get(UPFS_URL, {**params, "page_size": 200})
        return {item["id"] for item in response.data["results"]}

    def _titulares_exportados(self, cliente, params):
        response = cliente.get(UPFS_EXPORT_URL, params)
        assert response.status_code == status.HTTP_200_OK, response.content
        linhas = _linhas_csv(response.content)
        coluna = linhas[0].index("Titular")
        return {linha[coluna] for linha in linhas[1:]}

    @pytest.fixture
    def cenario_upfs(self, municipio_rn, municipio_ce, projeto, outro_projeto):
        comunidade = ComunidadeFactory(municipio=municipio_rn)
        UPFFactory(
            municipio=municipio_rn, projeto=projeto, comunidade=comunidade,
            _titular_nome="Maria Ativa RN",
        )
        UPFFactory(municipio=municipio_ce, projeto=outro_projeto, _titular_nome="Maria Ativa CE")
        UPFFactory(
            municipio=municipio_rn, projeto=projeto, _titular_nome="Maria Inativa", ativo=False,
        )
        UPFFactory(municipio=municipio_rn, projeto=outro_projeto, _titular_nome="José Ativo")
        return {
            "municipio": str(municipio_rn.pk),
            "territorio": str(municipio_ce.territory_id),
            "projeto": str(projeto.pk),
            "comunidade": str(comunidade.pk),
        }

    @pytest.mark.parametrize("params", [
        {},
        {"q": "Maria"},
        {"municipio": "{municipio}"},
        {"territorio": "{territorio}"},
        {"projeto": "{projeto}"},
        {"comunidade": "{comunidade}"},
        {"projeto": "{projeto}", "ativo": ""},
        {"ativo": "false"},
        {"ativo": ""},
        {"cadastrado_de": "2000-01-01", "cadastrado_ate": "2100-01-01"},
    ])
    def test_paridade_com_a_listagem(self, cliente_ugp, cenario_upfs, params):
        params = {chave: valor.format(**cenario_upfs) for chave, valor in params.items()}

        listagem = self._ids_listagem(cliente_ugp, params)
        esperados = set(
            UPF.all_objects.filter(pk__in=listagem).values_list(
                "titular__nome_completo", flat=True
            )
        )

        assert esperados, "o cenário deveria produzir ao menos uma UPF para o filtro"
        assert self._titulares_exportados(cliente_ugp, params) == esperados

    @pytest.mark.parametrize("params", [{"ativo": "xyz"}, {"ativa": "xyz"}])
    def test_ativo_invalido_retorna_400_em_vez_de_exportar_tudo(self, cliente_ugp, municipio_rn, params):
        UPFFactory(municipio=municipio_rn, ativo=False)

        response = cliente_ugp.get(UPFS_EXPORT_URL, params)

        assert response.status_code == status.HTTP_400_BAD_REQUEST
        assert "ativo" in response.data

    def test_ativo_invalido_tambem_e_400_na_listagem(self, cliente_ugp):
        response = cliente_ugp.get(UPFS_URL, {"ativo": "xyz"})

        assert response.status_code == status.HTTP_400_BAD_REQUEST

    def test_colunas(self, cliente_ugp, municipio_rn):
        UPFFactory(municipio=municipio_rn, _titular_nome="Maria", cpf="12345678901")

        response = cliente_ugp.get(UPFS_EXPORT_URL)

        linhas = _linhas_csv(response.content)
        assert linhas[0] == [
            "Estado", "Município", "Território", "Comunidade", "Titular", "CPF", "Data de cadastro",
        ]
        registro = dict(zip(linhas[0], linhas[1]))
        assert registro["Estado"] == "RN"
        assert registro["Território"] == municipio_rn.territory.nome

    def test_cpf_completo_para_ugp_e_mascarado_para_adt(self, cliente_ugp, municipio_rn, usuario_adt_rn):
        UPFFactory(municipio=municipio_rn, cpf="12345678901")

        ugp = dict(zip(*_linhas_csv(cliente_ugp.get(UPFS_EXPORT_URL).content)[:2]))
        adt = dict(zip(*_linhas_csv(_cliente(usuario_adt_rn).get(UPFS_EXPORT_URL).content)[:2]))

        assert ugp["CPF"] == "12345678901"
        assert adt["CPF"] == "123.***.***-01"

    def test_cpf_completo_para_super_admin_e_mascarado_para_articulador(
        self, municipio_rn, usuario_super_admin, usuario_articulador_rn
    ):
        UPFFactory(municipio=municipio_rn, cpf="12345678901")

        super_admin = _cliente(usuario_super_admin).get(UPFS_EXPORT_URL)
        articulador = _cliente(usuario_articulador_rn).get(UPFS_EXPORT_URL)

        assert dict(zip(*_linhas_csv(super_admin.content)[:2]))["CPF"] == "12345678901"
        assert dict(zip(*_linhas_csv(articulador.content)[:2]))["CPF"] == "123.***.***-01"

    def test_sem_campos_sensiveis(self, cliente_ugp, municipio_rn):
        UPFFactory(municipio=municipio_rn)

        cabecalho = _linhas_csv(cliente_ugp.get(UPFS_EXPORT_URL).content)[0]

        assert not {"Cor/Raça", "Condições de saúde", "Saúde"} & set(cabecalho)

    def test_escopo_territorial(self, auth_client_adt_rn, municipio_rn, municipio_ce):
        UPFFactory(municipio=municipio_rn, _titular_nome="Do RN")
        UPFFactory(municipio=municipio_ce, _titular_nome="Do CE")

        assert self._titulares_exportados(auth_client_adt_rn, {}) == {"Do RN"}

    def test_parametro_desconhecido_retorna_400(self, cliente_ugp):
        response = cliente_ugp.get(UPFS_EXPORT_URL, {"situacao": "ativas", "municipo": "1"})

        assert response.status_code == status.HTTP_400_BAD_REQUEST
        assert response.data["code"] == "parametro_desconhecido"
        assert response.data["parametros"] == ["municipo", "situacao"]

    @pytest.mark.parametrize("valor", ["true", "false", ""])
    def test_ativa_e_aceito_como_sinonimo_de_ativo(self, cliente_ugp, municipio_rn, valor):
        UPFFactory(municipio=municipio_rn, _titular_nome="Ativa")
        UPFFactory(municipio=municipio_rn, _titular_nome="Inativa", ativo=False)

        assert self._titulares_exportados(cliente_ugp, {"ativa": valor}) == (
            self._titulares_exportados(cliente_ugp, {"ativo": valor})
        )

    def test_ativo_prevalece_sobre_ativa(self, cliente_ugp, municipio_rn):
        UPFFactory(municipio=municipio_rn, _titular_nome="Ativa")
        UPFFactory(municipio=municipio_rn, _titular_nome="Inativa", ativo=False)

        exportados = self._titulares_exportados(cliente_ugp, {"ativo": "false", "ativa": "true"})

        assert exportados == {"Inativa"}

    def test_filtro_com_valor_invalido_retorna_400(self, cliente_ugp):
        response = cliente_ugp.get(UPFS_EXPORT_URL, {"municipio": "abc"})

        assert response.status_code == status.HTTP_400_BAD_REQUEST

    def test_acima_do_limite_vira_exportacao_assincrona(
        self, cliente_ugp, municipio_rn, django_capture_on_commit_callbacks, monkeypatch
    ):
        monkeypatch.setattr(upf_export, "UPF_EXPORT_SYNC_LIMIT", 2)
        UPFFactory.create_batch(3, municipio=municipio_rn)

        with patch(DELAY) as delay:
            with django_capture_on_commit_callbacks(execute=True):
                response = cliente_ugp.get(UPFS_EXPORT_URL, {"municipio": municipio_rn.pk})

        assert response.status_code == status.HTTP_202_ACCEPTED
        assert response.data["tipo"] == "upfs"
        assert response.data["total_registros"] == 3
        assert response.data["filtros"] == {"municipio": str(municipio_rn.pk)}
        delay.assert_called_once_with(response.data["id"])

        exportacao_service.executar_exportacao(response.data["id"])
        download = cliente_ugp.get(f"{_detalhe(response.data['id'])}download/")
        assert len(_linhas_csv(download.content)) == 4

    def test_no_limite_ainda_e_download_direto(self, cliente_ugp, municipio_rn, monkeypatch):
        monkeypatch.setattr(upf_export, "UPF_EXPORT_SYNC_LIMIT", 2)
        UPFFactory.create_batch(2, municipio=municipio_rn)

        response = cliente_ugp.get(UPFS_EXPORT_URL)

        assert response.status_code == status.HTTP_200_OK
        assert ExportJob.objects.count() == 0
