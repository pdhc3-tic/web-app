import csv
import io
from datetime import datetime
from decimal import Decimal

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core.tests.factories import UserFactory
from apps.sgd.models import Demand
from apps.sgd.tests.factories import ApprovalStepFactory, DemandFactory, DemandRequestFactory
from apps.sgp.tests.factories import ActivityFactory, BudgetAllocationFactory, BudgetRubricaFactory, WorkPlanAcaoFactory

pytestmark = pytest.mark.django_db

LISTA = "/api/v1/sgd/demandas/"
EXPORTAR = "/api/v1/sgd/demandas/exportar/"
PAINEL = "/api/v1/sgd/painel/rubricas/"
CUSTO = "/api/v1/sgd/relatorios/custo-por-atividade/"


def _em(demanda, ano, mes, dia):
    # `criado_em` é auto_now_add: só `update` consegue recuar a data.
    Demand.objects.filter(pk=demanda.pk).update(
        criado_em=timezone.make_aware(datetime(ano, mes, dia, 12, 0)),
    )


def _ids(response):
    corpo = response.json()
    return {d["id"] for d in (corpo["results"] if isinstance(corpo, dict) else corpo)}


@pytest.fixture
def cenario(usuario_ugp, activity_rn, solicitante_rn, municipio_ce):
    """Três demandas conhecidas: duas no RN (submetida de jan/2026 e autorizada
    de mar/2026) e uma no CE (submetida de jan/2026)."""
    activity_ce = ActivityFactory(municipio=municipio_ce, acao=WorkPlanAcaoFactory())
    outro = UserFactory(email="outro@test.com")
    rn_sub = DemandFactory(activity=activity_rn, solicitante=solicitante_rn, status="submetida")
    rn_aut = DemandFactory(activity=activity_rn, solicitante=solicitante_rn, status="autorizada")
    ce_sub = DemandFactory(activity=activity_ce, solicitante=outro, status="submetida")
    _em(rn_sub, 2026, 1, 15)
    _em(rn_aut, 2026, 3, 10)
    _em(ce_sub, 2026, 1, 20)
    return {"rn_sub": rn_sub, "rn_aut": rn_aut, "ce_sub": ce_sub, "activity_ce": activity_ce}


def test_filtros_combinados_retornam_exatamente_o_conjunto(auth_client_ugp, cenario, territory_rn):
    response = auth_client_ugp.get(LISTA, {
        "status": "submetida", "territorio": territory_rn.pk,
        "periodo_inicio": "2026-01-01", "periodo_fim": "2026-01-31",
    })
    assert response.status_code == 200
    assert _ids(response) == {cenario["rn_sub"].pk}


def test_filtros_individuais(auth_client_ugp, cenario, rubrica_diarias, solicitante_rn, activity_rn):
    DemandRequestFactory(demanda=cenario["rn_aut"], tipo="diaria", rubrica=rubrica_diarias)
    DemandRequestFactory(demanda=cenario["ce_sub"], tipo="grafico")
    ugp = auth_client_ugp
    rn_aut, rn_sub, ce_sub = cenario["rn_aut"], cenario["rn_sub"], cenario["ce_sub"]

    assert _ids(ugp.get(LISTA, {"tipo": "diaria"})) == {rn_aut.pk}
    assert _ids(ugp.get(LISTA, {"rubrica": "diarias"})) == {rn_aut.pk}
    assert _ids(ugp.get(LISTA, {"solicitante": solicitante_rn.pk})) == {rn_aut.pk, rn_sub.pk}
    assert _ids(ugp.get(LISTA, {"atividade": cenario["activity_ce"].pk})) == {ce_sub.pk}
    assert _ids(ugp.get(LISTA, {"acao": activity_rn.acao_id})) == {rn_aut.pk, rn_sub.pk}
    assert _ids(ugp.get(LISTA, {"submeta": activity_rn.acao.submeta_id})) == {rn_aut.pk, rn_sub.pk}
    assert _ids(ugp.get(LISTA, {"meta": activity_rn.acao.meta_id})) == {rn_aut.pk, rn_sub.pk}
    assert _ids(ugp.get(LISTA, {"periodo_inicio": "2026-02-01"})) == {rn_aut.pk}


def test_filtro_responsavel_pela_etapa_atual(
    auth_client_ugp, cenario, usuario_articulador_rn, usuario_articulador_ce, usuario_fgd, solicitante_rn,
):
    rn_aut, rn_sub = cenario["rn_aut"], cenario["rn_sub"]
    pre = DemandFactory(
        activity=rn_sub.activity, solicitante=solicitante_rn, status="pre_autorizada",
    )

    def por(user):
        return _ids(auth_client_ugp.get(LISTA, {"responsavel_etapa_atual": user.pk}))

    assert por(usuario_articulador_rn) == {rn_sub.pk}
    assert por(usuario_articulador_ce) == {cenario["ce_sub"].pk}
    assert por(usuario_fgd) == {rn_aut.pk}
    assert por(solicitante_rn) == set()
    assert pre.pk not in por(usuario_fgd)


def test_filtro_invalido_devolve_400(auth_client_ugp, cenario):
    assert auth_client_ugp.get(LISTA, {"status": "inexistente"}).status_code == 400
    assert auth_client_ugp.get(LISTA, {"periodo_inicio": "ontem"}).status_code == 400


def test_listagem_e_paginada_e_respeita_rbac(usuario_ugp, solicitante_rn, cenario):
    # Clientes próprios: as fixtures auth_client_* compartilham o mesmo
    # `api_client`, e a última autenticação sobrescreve as anteriores.
    cliente_ugp = APIClient()
    cliente_ugp.force_authenticate(user=usuario_ugp)
    cliente_adt = APIClient()
    cliente_adt.force_authenticate(user=solicitante_rn)

    corpo = cliente_ugp.get(LISTA, {"limit": 2}).json()
    assert corpo["count"] == 3
    assert len(corpo["results"]) == 2
    # ADT só enxerga as próprias, mesmo sem filtro.
    assert _ids(cliente_adt.get(LISTA)) == {cenario["rn_sub"].pk, cenario["rn_aut"].pk}


def test_historico_em_ordem_cronologica_com_dados_de_cada_etapa(
    auth_client_ugp, cenario, usuario_articulador_rn, usuario_ugp,
):
    demanda = cenario["rn_sub"]
    segunda = ApprovalStepFactory(
        demanda=demanda, etapa="autorizacao", responsavel=usuario_ugp, acao="aprovado", justificativa="ok",
    )
    primeira = ApprovalStepFactory(
        demanda=demanda, etapa="pre_autorizacao", responsavel=usuario_articulador_rn,
        acao="devolvido", justificativa="ajustar valores",
    )
    # Imutável após criado: só `update` de queryset recua o timestamp.
    type(segunda).objects.filter(pk=segunda.pk).update(criado_em=timezone.make_aware(datetime(2026, 2, 2)))
    type(primeira).objects.filter(pk=primeira.pk).update(criado_em=timezone.make_aware(datetime(2026, 2, 1)))

    response = auth_client_ugp.get(f"{LISTA}{demanda.pk}/historico/")

    assert response.status_code == 200
    etapas = response.json()
    assert [e["id"] for e in etapas] == [primeira.pk, segunda.pk]
    assert etapas[0]["responsavel_id"] == usuario_articulador_rn.pk
    assert etapas[0]["responsavel_nome"] == usuario_articulador_rn.nome
    assert etapas[0]["justificativa"] == "ajustar valores"
    assert etapas[0]["etapa"] == "pre_autorizacao"
    assert etapas[0]["acao"] == "devolvido"
    assert etapas[0]["criado_em"]


def test_historico_de_demanda_fora_do_escopo_devolve_404(auth_client_solicitante, cenario):
    assert auth_client_solicitante.get(f"{LISTA}{cenario['ce_sub'].pk}/historico/").status_code == 404


def test_painel_por_rubrica_bate_com_valores_agregados(
    auth_client_ugp, cenario, rubrica_diarias, activity_rn,
):
    rn_aut, rn_sub = cenario["rn_aut"], cenario["rn_sub"]
    DemandRequestFactory(
        demanda=rn_aut, tipo="diaria", rubrica=rubrica_diarias,
        valor_estimado=1000, valor_autorizado=800,
    )
    DemandRequestFactory(
        demanda=rn_sub, tipo="diaria", rubrica=rubrica_diarias, valor_estimado=400,
    )
    # Rascunho e recusada não entram em "solicitado".
    DemandRequestFactory(
        demanda=DemandFactory(activity=activity_rn, status="rascunho"),
        tipo="diaria", rubrica=rubrica_diarias, valor_estimado=9999,
    )
    DemandRequestFactory(
        demanda=DemandFactory(activity=activity_rn, status="recusada"),
        tipo="diaria", rubrica=rubrica_diarias, valor_estimado=7777,
    )
    BudgetAllocationFactory(
        meta=activity_rn.acao.meta, rubrica=rubrica_diarias, nivel="nacional", territorio=None,
        valor_alocado=10_000, valor_comprometido=1_000, valor_executado=500,
    )

    response = auth_client_ugp.get(PAINEL)

    assert response.status_code == 200
    linha = next(r for r in response.json() if r["rubrica_slug"] == "diarias")
    assert Decimal(linha["solicitado"]) == Decimal("1400")
    assert Decimal(linha["autorizado"]) == Decimal("800")
    assert Decimal(linha["executado"]) == Decimal("500")
    assert Decimal(linha["disponivel"]) == Decimal("8500")


def test_painel_e_relatorio_so_para_perfis_de_gestao(auth_client_solicitante):
    assert auth_client_solicitante.get(PAINEL).status_code == 403
    assert auth_client_solicitante.get(CUSTO).status_code == 403


def test_relatorio_de_custo_soma_todas_as_demandas_da_atividade(
    auth_client_ugp, cenario, rubrica_diarias, activity_rn,
):
    DemandRequestFactory(
        demanda=cenario["rn_aut"], tipo="diaria", rubrica=rubrica_diarias,
        valor_estimado=1000, valor_autorizado=900, valor_pago=850,
    )
    DemandRequestFactory(
        demanda=cenario["rn_sub"], tipo="diaria", rubrica=rubrica_diarias, valor_estimado=300,
    )
    DemandRequestFactory(
        demanda=DemandFactory(activity=activity_rn, status="cancelada"),
        tipo="diaria", rubrica=rubrica_diarias, valor_estimado=5000,
    )
    DemandRequestFactory(demanda=cenario["ce_sub"], tipo="grafico", valor_estimado=200)

    response = auth_client_ugp.get(CUSTO)

    assert response.status_code == 200
    entregas = {e["acao_id"]: e for e in response.json()}
    assert len(entregas) == 2
    entrega_rn = entregas[activity_rn.acao_id]
    assert entrega_rn["demandas"] == 2
    assert Decimal(entrega_rn["estimado"]) == Decimal("1300")
    assert Decimal(entrega_rn["autorizado"]) == Decimal("900")
    assert Decimal(entrega_rn["pago"]) == Decimal("850")
    assert [a["atividade_id"] for a in entrega_rn["atividades"]] == [activity_rn.pk]
    assert Decimal(entrega_rn["atividades"][0]["estimado"]) == Decimal("1300")
    assert Decimal(entregas[cenario["activity_ce"].acao_id]["estimado"]) == Decimal("200")


def test_exportacao_csv_reflete_os_filtros_aplicados(auth_client_ugp, cenario, rubrica_diarias):
    DemandRequestFactory(
        demanda=cenario["rn_sub"], tipo="diaria", rubrica=rubrica_diarias, valor_estimado=250,
    )

    territorio_id = cenario["rn_sub"].activity.municipio.territory_id

    response = auth_client_ugp.get(
        EXPORTAR, {"formato": "csv", "status": "submetida", "territorio": territorio_id},
    )

    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    linhas = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    cabecalho, dados = linhas[0], linhas[1:]
    assert len(dados) == 1
    registro = dict(zip(cabecalho, dados[0]))
    assert registro["ID"] == str(cenario["rn_sub"].pk)
    assert registro["Valor estimado (R$)"] == "250.00"
    assert registro["Rubricas"] == "Diárias"


def test_exportacao_xlsx_reflete_os_filtros_aplicados(auth_client_ugp, cenario):
    response = auth_client_ugp.get(EXPORTAR, {"formato": "xlsx", "periodo_inicio": "2026-02-01"})

    assert response.status_code == 200
    planilha = load_workbook(io.BytesIO(response.content)).active
    linhas = list(planilha.iter_rows(values_only=True))
    assert [str(linha[0]) for linha in linhas[1:]] == [str(cenario["rn_aut"].pk)]


def test_exportacao_sem_filtro_traz_tudo_e_filtro_invalido_falha(auth_client_ugp, cenario):
    response = auth_client_ugp.get(EXPORTAR)
    assert len(response.content.decode("utf-8-sig").strip().splitlines()) == 4
    assert auth_client_ugp.get(EXPORTAR, {"status": "xyz"}).status_code == 400
    assert auth_client_ugp.get(EXPORTAR, {"formato": "pdf"}).status_code == 400
