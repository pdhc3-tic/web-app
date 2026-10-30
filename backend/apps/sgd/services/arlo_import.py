"""Importação do retorno de pagamento do Arlo (SGD-RF32, RF24).

Cada linha da planilha é processada numa transação própria: uma linha com
erro vira entrada em `erros_json` e não impede as válidas do mesmo arquivo.
"""
import csv
import logging
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from urllib.parse import urlparse

from django.db import transaction

from apps.core.signals.audit import clear_audit_context, set_audit_context
from apps.core.storage import StorageObjectNotFound, get_storage
from apps.sgd.exceptions import ErroDeDominio, TransicaoInvalida
from apps.sgd.models.arlo_import import ArloImport
from apps.sgd.models.demand import Demand
from apps.sgd.models.demand_document import DemandDocument
from apps.sgd.services import balance as balance_service
from apps.sgd.services import workflow as workflow_service
from apps.sgd.services.arlo_mapping import CAMPOS_IMPORTACAO_OBRIGATORIOS, obter_mapeamento

logger = logging.getLogger(__name__)

FORMATOS_DATA = ("%Y-%m-%d", "%d/%m/%Y")
MENSAGEM_ARQUIVO_NAO_ENCONTRADO = (
    "O arquivo enviado não foi encontrado no armazenamento. Envie a planilha novamente."
)


class LinhaInvalida(Exception):
    def __init__(self, erro: str, campo: str | None = None):
        super().__init__(erro)
        self.erro = erro
        self.campo = campo


class ArquivoInvalido(Exception):
    pass


def _ler_linhas(conteudo: bytes, nome: str) -> tuple[list[str], list[dict]]:
    if nome.lower().endswith(".xlsx"):
        return _ler_xlsx(conteudo)
    texto = conteudo.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(texto[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(StringIO(texto), dialect=dialect)
    return list(reader.fieldnames or []), [dict(linha) for linha in reader]


def _ler_xlsx(conteudo: bytes) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook

    planilha = load_workbook(BytesIO(conteudo), read_only=True, data_only=True).active
    linhas = planilha.iter_rows(values_only=True)
    cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas, [])]
    return cabecalho, [
        {cabecalho[i]: ("" if v is None else v) for i, v in enumerate(linha) if i < len(cabecalho)}
        for linha in linhas
    ]


def _valor(linha: dict, colunas: dict, campo: str):
    coluna = colunas.get(campo)
    bruto = linha.get(coluna) if coluna else None
    if isinstance(bruto, str):
        bruto = bruto.strip()
    return "" if bruto is None else bruto


def _parse_valor(bruto) -> Decimal:
    if isinstance(bruto, (int, float, Decimal)):
        texto = str(bruto)
    else:
        texto = str(bruto).replace("R$", "").strip()
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
    try:
        valor = Decimal(texto)
    except InvalidOperation:
        raise LinhaInvalida(f"Valor pago em formato inválido: '{bruto}'.", "valor_pago")
    if not valor.is_finite() or valor < 0:
        raise LinhaInvalida(f"Valor pago inválido: '{bruto}'.", "valor_pago")
    return valor.quantize(Decimal("0.01"))


def _parse_data(bruto) -> date:
    if isinstance(bruto, datetime):
        return bruto.date()
    if isinstance(bruto, date):
        return bruto
    for formato in FORMATOS_DATA:
        try:
            return datetime.strptime(str(bruto), formato).date()
        except ValueError:
            continue
    raise LinhaInvalida(f"Data do pagamento em formato inválido: '{bruto}'.", "data_pagamento")


def _parse_inteiro(bruto, campo: str) -> int:
    try:
        return int(Decimal(str(bruto)))
    except (InvalidOperation, ValueError):
        raise LinhaInvalida(f"'{bruto}' não é um identificador válido.", campo)


def _validar_comprovante(bruto) -> str:
    if not bruto:
        return ""
    parsed = urlparse(str(bruto))
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise LinhaInvalida("Comprovante deve ser uma URL http(s).", "comprovante")
    return str(bruto)


def _exigir_pagavel(demand) -> None:
    if demand.status not in workflow_service.STATUS_PAGAVEIS_PELO_ARLO:
        raise LinhaInvalida(
            f"Demanda #{demand.pk} está '{demand.status}' — só autorizadas/em atendimento recebem pagamento.",
            "demanda_id",
        )


def _localizar_solicitacao(demand, bruto_solicitacao):
    solicitacoes = list(demand.solicitacoes.select_related("rubrica").order_by("rubrica_id", "pk"))
    if bruto_solicitacao != "":
        pk = _parse_inteiro(bruto_solicitacao, "solicitacao_id")
        for solicitacao in solicitacoes:
            if solicitacao.pk == pk:
                return solicitacao, solicitacoes
        raise LinhaInvalida(f"Solicitação #{pk} não pertence à demanda #{demand.pk}.", "solicitacao_id")
    if len(solicitacoes) != 1:
        raise LinhaInvalida(
            f"Demanda #{demand.pk} tem {len(solicitacoes)} solicitações — informe o ID da solicitação.",
            "solicitacao_id",
        )
    return solicitacoes[0], solicitacoes


def _processar_linha(importacao, linha: dict, colunas: dict) -> dict | None:
    """Aplica o pagamento da linha e devolve a divergência entre pago e
    autorizado, ou None quando os valores batem."""
    for campo in CAMPOS_IMPORTACAO_OBRIGATORIOS:
        if _valor(linha, colunas, campo) == "":
            raise LinhaInvalida(f"Campo obrigatório ausente: {colunas[campo]}.", campo)

    demanda_id = _parse_inteiro(_valor(linha, colunas, "demanda_id"), "demanda_id")
    data_pagamento = _parse_data(_valor(linha, colunas, "data_pagamento"))
    valor_pago = _parse_valor(_valor(linha, colunas, "valor_pago"))
    comprovante = _validar_comprovante(_valor(linha, colunas, "comprovante"))
    numero_processo = str(_valor(linha, colunas, "numero_processo"))[:64]

    demand = Demand.objects.select_related("activity").filter(pk=demanda_id).first()
    if demand is None:
        raise LinhaInvalida(f"Demanda #{demanda_id} inexistente.", "demanda_id")
    _exigir_pagavel(demand)

    solicitacao, _ = _localizar_solicitacao(demand, _valor(linha, colunas, "solicitacao_id"))
    if solicitacao.valor_pago is not None:
        raise LinhaInvalida(f"Solicitação #{solicitacao.pk} já tem pagamento registrado.", "solicitacao_id")

    usuario = importacao.operado_por
    reservado = workflow_service.valor_autorizado(solicitacao)
    # O saldo antes da Demand, na ordem de lock de backend/docs/orcamento.md.
    try:
        workflow_service.registrar_pagamento(
            solicitacao, valor_pago=valor_pago, origem=workflow_service.ORIGEM_ARLO, usuario=usuario,
            importacao=importacao, numero_processo=numero_processo, data_pagamento=data_pagamento.isoformat(),
        )
    except ErroDeDominio as exc:
        raise LinhaInvalida(f"Falha ao executar no orçamento: {exc.mensagem}")
    # Relido sob o lock: outra importação pode ter iniciado o atendimento, ou a
    # demanda pode ter sido cancelada ou concluída, enquanto o saldo era executado.
    demand.status = workflow_service.travar_e_ler_status(demand)
    _exigir_pagavel(demand)
    if demand.status == "autorizada":
        workflow_service.executar_transicao(
            demand, "em_atendimento", ator=usuario, origem=workflow_service.ORIGEM_ARLO,
        )

    balance_service.gravar_estado_da_reserva([solicitacao], outros_campos=("valor_pago",))
    demand.recalcular_totais()
    demand.numero_processo = numero_processo
    demand.save(update_fields=["numero_processo"])

    if comprovante:
        DemandDocument.objects.create(
            demanda=demand, arquivo_key="", arquivo_url=comprovante, tipo="comprovante",
            nome_original=urlparse(comprovante).path.rsplit("/", 1)[-1] or "comprovante",
            enviado_por=usuario,
        )

    if not demand.solicitacoes.filter(valor_pago__isnull=True).exists():
        workflow_service.executar_transicao(
            demand, "concluida", ator=usuario, origem=workflow_service.ORIGEM_ARLO,
            justificativa=(
                f"Pagamento em {data_pagamento:%d/%m/%Y} — processo {numero_processo} "
                f"(importação Arlo #{importacao.pk})."
            ),
        )

    if valor_pago == reservado:
        return None
    diferenca = valor_pago - reservado
    return {
        "demanda_id": demand.pk, "demanda_titulo": demand.titulo, "solicitacao_id": solicitacao.pk,
        "valor_autorizado": str(reservado), "valor_pago": str(valor_pago), "diferenca": str(diferenca),
        "risco_glosa": diferenca > 0,
    }


def processar_importacao(importacao: ArloImport, *, transacao_da_linha=transaction.atomic) -> ArloImport:
    """Lê o arquivo de retorno de `importacao` e aplica linha a linha, cada linha
    numa transação aberta por `transacao_da_linha` (a task passa uma que define o
    papel da sessão para o RLS). Quem chama não abre transação em volta: com o
    arquivo inteiro numa só, uma linha seguraria alocação e `Demand` enquanto a
    seguinte trava o limite de outro solicitante, fora da ordem de lock.

    Roda fora de um request, então abre o contexto de auditoria com o usuário e
    o IP de quem enviou o arquivo — é dele que as movimentações de saldo e as
    transições de status herdam usuário/IP (SGD RNF de rastreabilidade)."""
    set_audit_context(user=importacao.operado_por, ip=importacao.ip_origem, user_agent="arlo-import")
    try:
        return _processar_importacao(importacao, transacao_da_linha)
    finally:
        clear_audit_context()


def _processar_importacao(importacao: ArloImport, transacao_da_linha) -> ArloImport:
    importacao.status = ArloImport.Status.PROCESSANDO
    importacao.save(update_fields=["status"])

    try:
        conteudo = get_storage().read_bytes(importacao.arquivo_key)
        cabecalho, linhas = _ler_linhas(conteudo, importacao.nome_original or importacao.arquivo_key)
        mapeamento = obter_mapeamento()
        importacao.mapeamento_snapshot = mapeamento
        importacao.save(update_fields=["mapeamento_snapshot"])
        colunas = {campo: item["coluna"] for campo, item in mapeamento["importacao"].items()}
        ausentes = [coluna for campo, coluna in colunas.items()
                    if campo in CAMPOS_IMPORTACAO_OBRIGATORIOS and coluna not in cabecalho]
        if ausentes:
            raise ArquivoInvalido(f"Colunas obrigatórias ausentes na planilha: {ausentes}.")
    except StorageObjectNotFound:
        logger.warning("Arquivo da importação Arlo #%s não encontrado: %s.", importacao.pk, importacao.arquivo_key)
        return _falhar(importacao, MENSAGEM_ARQUIVO_NAO_ENCONTRADO)
    except (ArquivoInvalido, UnicodeDecodeError, ValueError, KeyError) as exc:
        return _falhar(importacao, str(exc) or "Arquivo ilegível.")

    erros, divergencias, ok, total = [], [], 0, 0
    for numero, linha in enumerate(linhas, start=2):  # linha 1 é o cabeçalho
        if not any(str(v).strip() for v in linha.values() if v is not None):
            continue
        total += 1
        try:
            with transacao_da_linha():
                divergencia = _processar_linha(importacao, linha, colunas)
                # O progresso entra na transação da linha: se o processo cair no meio
                # do arquivo, a importação mostra exatamente as linhas já aplicadas.
                _gravar_progresso(
                    importacao, total=total, ok=ok + 1, erros=erros,
                    divergencias=[*divergencias, divergencia] if divergencia else divergencias,
                )
            ok += 1
            if divergencia:
                divergencias.append(divergencia)
            continue
        except LinhaInvalida as exc:
            erros.append({"linha": numero, "campo": exc.campo, "erro": exc.erro})
        except TransicaoInvalida as exc:
            erros.append({"linha": numero, "campo": None, "erro": exc.mensagem})
        except Exception:
            logger.exception("Erro inesperado na linha %s da importação Arlo #%s.", numero, importacao.pk)
            erros.append({"linha": numero, "campo": None, "erro": "Erro inesperado ao processar a linha."})
        _gravar_progresso(importacao, total=total, ok=ok, erros=erros, divergencias=divergencias)

    importacao.total_registros = total
    importacao.registros_ok = ok
    importacao.erros_json = erros
    importacao.divergencias_json = divergencias
    importacao.status = ArloImport.Status.CONCLUIDO
    importacao.save(update_fields=["total_registros", "registros_ok", "erros_json", "divergencias_json", "status"])
    return importacao


def _gravar_progresso(importacao: ArloImport, *, total: int, ok: int, erros: list, divergencias: list) -> None:
    ArloImport.objects.filter(pk=importacao.pk).update(
        total_registros=total, registros_ok=ok, erros_json=erros, divergencias_json=divergencias,
    )


def _falhar(importacao: ArloImport, erro: str) -> ArloImport:
    importacao.status = ArloImport.Status.FALHOU
    importacao.erros_json = [{"linha": 1, "erro": erro}]
    importacao.save(update_fields=["status", "erros_json"])
    return importacao
